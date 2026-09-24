import os
from io import BytesIO
from datetime import date

import streamlit as st
from dotenv import load_dotenv
from google import genai
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

load_dotenv()

st.set_page_config(page_title="LegalEase AI", page_icon="⚖️", layout="wide")
st.title("⚖️ LegalEase AI")
st.caption("AI-assisted legal document drafting, editing and downloading.")

st.info("This application creates drafts for educational/general-information use. It is not legal advice. Have an appropriately qualified lawyer review documents before real-world use.")

def get_client():
    key = os.getenv("GEMINI_API_KEY", "").strip()
    if not key or key == "PASTE_YOUR_API_KEY_HERE":
        return None
    return genai.Client(api_key=key)

def generate_with_gemini(doc_type, parties, terms, effective_date):
    client = get_client()
    if client is None:
        raise ValueError("GEMINI_API_KEY is missing. Add it to the .env file and restart the app.")

    prompt = f"""
You are a legal-document drafting assistant.
Create a clear, structured FIRST DRAFT of a {doc_type}.

Inputs:
Parties: {parties}
Effective date: {effective_date}
Terms and conditions: {terms}

Requirements:
- Use headings and numbered clauses.
- Do not invent names, addresses, prices, dates or obligations that were not provided.
- If important information is missing, mark it as [TO BE COMPLETED].
- Use neutral professional language.
- Include signature sections where appropriate.
- Add a short "Drafting Notice" at the end saying the document should be reviewed for applicable local law.
- This is drafting assistance, not legal advice.
Return only the document text.
"""
    response = client.interactions.create(
        model="gemini-3.8-flash",
        input=prompt
    )
    return response.output_text

def make_docx(text, title):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    p = doc.add_paragraph()
    r = p.add_run(title.upper())
    r.bold = True
    r.font.size = Pt(16)
    for line in text.splitlines():
        doc.add_paragraph(line)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()

def make_pdf(text, title):
    out = BytesIO()
    pdf = SimpleDocTemplate(out, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    story = [Paragraph(title.upper(), styles["Title"]), Spacer(1, 12)]
    for line in text.splitlines():
        safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if not safe.strip():
            story.append(Spacer(1, 8))
        elif safe.isupper() and len(safe) < 70:
            story.append(Paragraph(f"<b>{safe}</b>", styles["Heading2"]))
        else:
            story.append(Paragraph(safe, styles["BodyText"]))
    pdf.build(story)
    return out.getvalue()

with st.sidebar:
    st.header("Document Details")
    doc_type = st.selectbox(
        "Document type",
        ["Agreement", "NDA (Non-Disclosure Agreement)", "Lease Agreement",
         "Employment Contract", "Employment Offer Letter", "Freelance Contract"]
    )
    parties = st.text_area(
        "Parties involved",
        placeholder="Example: Jane Doe (Service Provider), ABC Ltd (Client)"
    )
    terms = st.text_area(
        "Terms & conditions",
        placeholder="Example: Payment within 30 days; Confidentiality; 15 days notice for termination"
    )
    effective_date = st.date_input("Effective date", value=date.today())
    generate = st.button("✨ Generate with Gemini", type="primary", use_container_width=True)

if "document" not in st.session_state:
    st.session_state.document = ""

if generate:
    if not parties.strip() or not terms.strip():
        st.error("Please enter the parties and terms.")
    else:
        with st.spinner("Generating draft with Gemini..."):
            try:
                st.session_state.document = generate_with_gemini(
                    doc_type, parties, terms, effective_date.strftime("%d/%m/%Y")
                )
                st.success("Draft generated. You can edit it below.")
            except Exception as e:
                st.error(f"Generation failed: {e}")

st.subheader("Editable Document")
st.session_state.document = st.text_area(
    "Edit the generated document",
    value=st.session_state.document,
    height=560,
    placeholder="Generate a document, then edit the text here."
)

if st.session_state.document.strip():
    text = st.session_state.document
    st.subheader("Download")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.download_button("⬇️ Download TXT", text.encode("utf-8"),
                           "legalease_document.txt", "text/plain", use_container_width=True)
    with c2:
        st.download_button("⬇️ Download DOCX", make_docx(text, doc_type),
                           "legalease_document.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           use_container_width=True)
    with c3:
        st.download_button("⬇️ Download PDF", make_pdf(text, doc_type),
                           "legalease_document.pdf", "application/pdf",
                           use_container_width=True)

st.divider()
st.caption("LegalEase AI • Python + Streamlit + Google Gemini")
